"""Text extraction for uploaded files; the extracted text is what gets chunked and embedded."""

from io import BytesIO
from pathlib import Path

import docx
import pypdf

from app.models import MaterialType

SUPPORTED_EXTENSIONS = {
    ".pdf": MaterialType.pdf,
    ".docx": MaterialType.docx,
    ".txt": MaterialType.text_file,
    ".md": MaterialType.text_file,
}


class ExtractionError(Exception):
    pass


def material_type_for(filename: str) -> MaterialType | None:
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower())


def extract_text(data: bytes, material_type: MaterialType) -> str:
    try:
        if material_type == MaterialType.pdf:
            text = _pdf_text(data)
        elif material_type == MaterialType.docx:
            text = _docx_text(data)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except Exception as exc:  # parsers raise many different errors for corrupt files
        raise ExtractionError("The file couldn't be read. Is it corrupted?") from exc

    text = text.replace("\x00", "").strip()  # Postgres text columns can't store NUL
    if not text:
        if material_type == MaterialType.pdf:
            raise ExtractionError(
                "No text found in this PDF. Scanned PDFs without a text layer aren't supported."
            )
        raise ExtractionError("This file doesn't contain any text.")
    return text


def _pdf_text(data: bytes) -> str:
    reader = pypdf.PdfReader(BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    document = docx.Document(BytesIO(data))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(block for block in blocks if block.strip())
